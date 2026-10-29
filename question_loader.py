import os
import openpyxl

def parse_choice_answer_index(cell_val, choices):
    """
    H列のセル指定 (C, D, E, F, C列, D列, 1, 2, 3, 4 等) やテキストから
    0-based インデックス (0, 1, 2, 3) を特定する
    """
    if cell_val is None:
        return 0
    val_str = str(cell_val).strip()
    val_upper = val_str.upper()
    
    col_map = {
        "C": 0, "D": 1, "E": 2, "F": 3,
        "C列": 0, "D列": 1, "E列": 2, "F列": 3,
        "1": 0, "2": 1, "3": 2, "4": 3,
        "0": 0
    }
    if val_upper in col_map:
        return col_map[val_upper]
    
    if len(val_upper) >= 2 and val_upper[0] in col_map:
        return col_map[val_upper[0]]
        
    for idx, c in enumerate(choices):
        if str(c).strip().lower() == val_str.lower():
            return idx
            
    return 0

def load_questions(filepath: str, difficulty: str = "all") -> list:
    """
    Excelファイル(question.xlsx)から問題を読み込む
    - 5段階難易度: easy, normal-, normal, normal+, difficult
    """
    if not os.path.exists(filepath):
        print(f"[Warning] Questions file not found at {filepath}")
        return []

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        print(f"[Warning] Failed to open Excel file: {e}")
        return []

    all_levels = ["easy", "normal-", "normal", "normal+", "difficult"]
    target_sheets = []

    if difficulty in all_levels:
        if difficulty in wb.sheetnames:
            target_sheets = [difficulty]
    else:
        for sheet_name in all_levels:
            if sheet_name in wb.sheetnames:
                target_sheets.append(sheet_name)
        if not target_sheets:
            target_sheets = wb.sheetnames

    questions = []
    global_id = 1

    for sheet_name in target_sheets:
        ws = wb[sheet_name]
        for row in list(ws.iter_rows(min_row=2, values_only=True)):
            if not row or len(row) < 1:
                continue
            
            q_text = row[0]  # A列: 問題文
            if not q_text or str(q_text).strip() == "":
                continue
            
            q_type = str(row[1]).strip().lower() if len(row) > 1 and row[1] is not None else "text"  # B列: 問題形式
            if q_type not in ("mhy", "choice", "text"):
                q_type = "text"

            item = {
                "id": global_id,
                "sheet": sheet_name,
                "difficulty": sheet_name,
                "question": str(q_text).strip(),
                "type": q_type
            }
            global_id += 1

            if q_type == "choice":
                choices = []
                for col_i in range(2, 6):  # C～F列
                    val = row[col_i] if len(row) > col_i and row[col_i] is not None else ""
                    choices.append(str(val).strip())
                
                if not any(choices):
                    choices = ["選択肢1", "選択肢2", "選択肢3", "選択肢4"]

                h_val = row[7] if len(row) > 7 and row[7] is not None else "C"  # H列: 解答セル指定
                ans_idx = parse_choice_answer_index(h_val, choices)

                item["choices"] = choices
                item["answer_index"] = ans_idx

            elif q_type == "mhy":
                g_val = row[6] if len(row) > 6 and row[6] is not None else ""  # G列: ダミー文字列
                if g_val:
                    item["dummy_chars"] = str(g_val).strip()

                h_val = row[7] if len(row) > 7 and row[7] is not None else ""  # H列: 模範解答1つ
                ans = str(h_val).strip()
                item["answers"] = [ans] if ans else ["正解"]

            elif q_type == "text":
                h_val = row[7] if len(row) > 7 and row[7] is not None else ""  # H列: カンマ区切り模範解答複数
                val_str = str(h_val).replace("，", ",")
                raw_answers = [a.strip() for a in val_str.split(",") if a.strip()]
                item["answers"] = raw_answers if raw_answers else ["正解"]

            questions.append(item)

    return questions